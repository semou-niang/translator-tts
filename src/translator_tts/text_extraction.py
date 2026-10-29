"""Extraction locale bornée des quatre formats, sans persistance ni réseau."""

import csv
import io
import logging
import warnings
import zipfile
from pathlib import PurePath
from xml.etree.ElementTree import ParseError

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from translator_tts.config import Settings
from translator_tts.utils import ExtractionError, InputValidationError, validate_source_text

MAX_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_CELLS = 100000
MAX_PAGES = 200

# Les diagnostics du parseur peuvent contenir des fragments du document malformé.
logging.getLogger("pypdf").setLevel(logging.CRITICAL)


def _decode(data: bytes) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("cp1252")
        except UnicodeDecodeError:
            raise ExtractionError("Encodage texte illisible.") from None
    if any(ord(char) < 32 and char not in "\n\r\t" for char in text):
        raise ExtractionError("Le fichier contient des données binaires.")
    if data.startswith((b"%PDF-", b"PK\x03\x04")):
        raise ExtractionError("Le contenu ne correspond pas au format texte annoncé.")
    return text


def _csv(data: bytes) -> str:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    try:
        for row in csv.reader(io.StringIO(text, newline=""), dialect, strict=True):
            if any(cell.strip() for cell in row):
                rows.append(" | ".join(cell for cell in row if cell.strip()))
    except csv.Error:
        raise ExtractionError("Le CSV est malformé ou contient des champs trop longs.") from None
    return "\n".join(rows)


def _pdf(data: bytes) -> str:
    if not data.startswith(b"%PDF-"):
        raise ExtractionError("Signature PDF invalide.")
    try:
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise ExtractionError("Les PDF chiffrés ne sont pas pris en charge.")
        if len(reader.pages) > MAX_PAGES:
            raise ExtractionError(f"Le PDF dépasse la limite de {MAX_PAGES} pages.")
        pages = []
        expanded = 0
        for page in reader.pages:
            contents = page.get_contents()
            if contents is not None:
                expanded += len(contents.get_data())
                if expanded > MAX_EXPANDED_BYTES:
                    raise ExtractionError("Le contenu décompressé du PDF est trop volumineux.")
            pages.append(page.extract_text() or "")
        text = "\n\n".join(pages)
    except (
        PdfReadError,
        ValueError,
        TypeError,
        KeyError,
        OSError,
        RecursionError,
        NotImplementedError,
    ):
        raise ExtractionError("Impossible de lire ce PDF endommagé.") from None
    if not text.strip():
        raise ExtractionError(
            "Aucun texte exploitable ; les PDF scannés ne sont pas pris en charge."
        )
    return text


def _xlsx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            names = {item.filename for item in entries}
            if (
                len(entries) > 2000
                or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES
                or not {"[Content_Types].xml", "xl/workbook.xml"} <= names
                or any("vbaproject" in name.casefold() for name in names)
            ):
                raise ExtractionError("Structure XLSX invalide ou trop volumineuse.")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=True, keep_links=False
            )
        try:
            sections = []
            visited = 0
            for sheet in workbook:
                # Les dimensions annoncées peuvent être fausses ; le budget reste appliqué.
                if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_CELLS:
                    raise ExtractionError("La feuille Excel dépasse la limite de cellules.")
                sheet.reset_dimensions()
                rows = []
                for row in sheet.iter_rows():
                    visited += len(row)
                    if visited > MAX_CELLS:
                        raise ExtractionError("Le classeur dépasse la limite de cellules.")
                    values = [str(cell.value) for cell in row if cell.value is not None]
                    if any(value.strip() for value in values):
                        rows.append(" | ".join(values))
                if rows:
                    sections.append(f"Feuille : {sheet.title}\n" + "\n".join(rows))
            return "\n\n".join(sections)
        finally:
            workbook.close()
    except (
        zipfile.BadZipFile,
        InvalidFileException,
        ParseError,
        KeyError,
        ValueError,
        OSError,
        RuntimeError,
        NotImplementedError,
    ):
        raise ExtractionError("Impossible de lire ce classeur XLSX endommagé.") from None


def extract_text(file_name: str, file_bytes: bytes, settings: Settings | None = None) -> str:
    settings = settings or Settings()
    if len(file_bytes) > settings.max_file_bytes:
        raise ExtractionError(f"Le fichier dépasse {settings.max_file_bytes} octets.")
    extension = PurePath(file_name).suffix.lower()
    parsers = {".txt": _decode, ".csv": _csv, ".pdf": _pdf, ".xlsx": _xlsx}
    if extension not in parsers:
        raise ExtractionError("Format non pris en charge : TXT, PDF, CSV ou XLSX.")
    try:
        return validate_source_text(parsers[extension](file_bytes), settings.max_source_chars)
    except InputValidationError as error:
        raise ExtractionError(str(error)) from None

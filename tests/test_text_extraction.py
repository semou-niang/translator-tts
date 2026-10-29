"""UT-03..12 : fichiers synthétiques réellement extraits."""

import io
import zipfile

import openpyxl
import pytest
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from translator_tts.config import Settings
from translator_tts.text_extraction import extract_text
from translator_tts.utils import ExtractionError, InputValidationError, validate_source_text


def pdf_bytes(texts=(), encrypted=False):
    writer = PdfWriter()
    for text in texts or ("",):
        page = writer.add_blank_page(200, 200)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 10 10 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("synthetic-password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def workbook_bytes():
    workbook = openpyxl.Workbook()
    workbook.active.title = "Premier"
    workbook.active.append(["Bonjour", None, 42])
    workbook.create_sheet("Vide")
    workbook.create_sheet("Deuxième").append(["été", "=1+1"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "cp1252"])
def test_text_encodings(encoding):
    assert extract_text("../../texte.txt", "été\r\nBonjour".encode(encoding)) == "été\nBonjour"


@pytest.mark.parametrize("delimiter", [",", ";", "\t"])
def test_csv(delimiter):
    data = f'"été{delimiter}ici"{delimiter}42\nbonjour{delimiter}monde'
    assert extract_text("a.csv", data.encode()) == f"été{delimiter}ici | 42\nbonjour | monde"


def test_pdf_order():
    text = extract_text("a.pdf", pdf_bytes(("Hello", "World")))
    assert text.index("Hello") < text.index("World")


def test_workbook():
    assert extract_text("a.xlsx", workbook_bytes()) == (
        "Feuille : Premier\nBonjour | 42\n\nFeuille : Deuxième\nété"
    )


@pytest.mark.parametrize(
    "name,data",
    [
        ("a.txt", b""),
        ("a.csv", b"  "),
        ("a.txt", b"\x00abc"),
        ("a.txt", b"\x81"),
        ("a.csv", b'"unterminated'),
        ("a.pdf", b"fake"),
        ("a.pdf", b"%PDF-1.7\ntruncated"),
        ("a.xlsx", b"PKfake"),
        ("a.exe", b"hello"),
        ("a.txt", b"%PDF-1.7"),
        ("a.txt", b"PK\x03\x04"),
    ],
)
def test_invalid_files(name, data):
    with pytest.raises(ExtractionError):
        extract_text(name, data)


def test_pdf_empty_and_encrypted():
    for data in (pdf_bytes(), pdf_bytes(("Hello",), encrypted=True)):
        with pytest.raises(ExtractionError):
            extract_text("a.pdf", data)


def test_zip_budget():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", b"x" * (21 * 1024 * 1024))
        archive.writestr("[Content_Types].xml", "x")
    with pytest.raises(ExtractionError):
        extract_text("a.xlsx", output.getvalue())


def test_file_and_character_limits():
    with pytest.raises(ExtractionError):
        extract_text("a.txt", b"1234", Settings(max_file_bytes=3))
    assert validate_source_text("x" * 10000) == "x" * 10000
    for value in (" ", "x" * 10001):
        with pytest.raises(InputValidationError):
            validate_source_text(value)
        with pytest.raises(ExtractionError):
            extract_text("a.txt", value.encode())


@pytest.mark.parametrize("entry", ["xl/vbaProject.bin", "xl/vbaproject.bin", "xl/VBAPROJECT.BIN"])
def test_macro_archive_rejected_case_insensitively(entry):
    output = io.BytesIO(workbook_bytes())
    with zipfile.ZipFile(output, "a") as archive:
        archive.writestr(entry, b"synthetic-inert-macro-marker")
    with pytest.raises(ExtractionError, match="Structure XLSX"):
        extract_text("synthetic.xlsx", output.getvalue())
